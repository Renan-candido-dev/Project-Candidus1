"""Gera as planilhas funcionais das demonstrações práticas (dados fictícios).

Uso: python3 conteudo/planilhas/gerar_planilhas.py   (requer openpyxl)
Os arquivos .xlsx ficam nesta mesma pasta. As fórmulas usam nomes em inglês (padrão do
formato .xlsx); o Excel e o Google Planilhas em português exibem SOMA, SOMASE etc.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent
GREEN = "0E3B2E"
GOLD = "C9A24A"
IVORY = "F4F1EA"
CORAL_BG = "F6DDD6"
MONEY = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
PCT = "0.0%"
thin = Side(style="thin", color="DDDDDD")


def header(ws, row, cols):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.alignment = Alignment(horizontal="center", vertical="center")


def title(ws, text, sub="Dados fictícios, apenas para demonstração. Candidus Gestão Financeira."):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color=GREEN)
    ws["A2"] = sub
    ws["A2"].font = Font(italic=True, size=9, color="6E7C75")


def widths(ws, ws_widths):
    for col, w in ws_widths.items():
        ws.column_dimensions[col].width = w


def grid(ws, r1, r2, c1, c2):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(row=r, column=c).border = Border(top=thin, bottom=thin, left=thin, right=thin)


def label(ws, cell, text, bold=True, fill=IVORY):
    ws[cell] = text
    ws[cell].font = Font(bold=bold, color=GREEN)
    ws[cell].fill = PatternFill("solid", fgColor=fill)


def neg_red(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=CORAL_BG)))


# ------------------------------------------------------------------ V04
def v04():
    wb = Workbook()
    ws = wb.active
    ws.title = "Extrato"
    title(ws, "V04 · Extrato da empresa (junho): pessoal × empresa")
    header(ws, 4, ["Data", "Descrição", "Categoria", "Valor"])
    rows = [
        ("01/06", "Vendas no cartão", "Receitas", 12400),
        ("03/06", "Fornecedor Alfa", "Fornecedores", -3480),
        ("04/06", "Supermercado", "Gasto pessoal", -612),
        ("05/06", "Aluguel da loja", "Despesas fixas", -4800),
        ("08/06", "Pix de clientes", "Receitas", 6900),
        ("10/06", "Escola dos filhos", "Gasto pessoal", -1450),
        ("12/06", "Folha de pagamento", "Folha", -7200),
        ("15/06", "Streaming e apps", "Gasto pessoal", -89),
        ("18/06", "Guia DAS", "Impostos", -1380),
        ("22/06", "Viagem de férias", "Gasto pessoal", -2300),
    ]
    for i, r in enumerate(rows, 5):
        for j, v in enumerate(r, 1):
            ws.cell(row=i, column=j, value=v)
        ws.cell(row=i, column=4).number_format = MONEY
    last = 4 + len(rows)
    grid(ws, 4, last, 1, 4)
    label(ws, f"B{last + 2}", "Resultado do mês (extrato)")
    ws[f"D{last + 2}"] = f"=SUM(D5:D{last})"
    label(ws, f"B{last + 3}", "Gastos pessoais pagos pela empresa")
    ws[f"D{last + 3}"] = f'=SUMIF(C5:C{last},"Gasto pessoal",D5:D{last})'
    label(ws, f"B{last + 4}", "Resultado da operação (sem pessoais)")
    ws[f"D{last + 4}"] = f"=D{last + 2}-D{last + 3}"
    for r in range(last + 2, last + 5):
        ws[f"D{r}"].number_format = MONEY
        ws[f"D{r}"].font = Font(bold=True)
    neg_red(ws, f"D5:D{last + 4}")
    # lista de categorias + validação
    cats = wb.create_sheet("Categorias")
    title(cats, "Categorias (edite à vontade)")
    header(cats, 4, ["Categoria", "Total no mês"])
    names = ["Receitas", "Fornecedores", "Folha", "Despesas fixas", "Impostos", "Gasto pessoal"]
    for i, n in enumerate(names, 5):
        cats.cell(row=i, column=1, value=n)
        cats.cell(row=i, column=2, value=f"=SUMIF(Extrato!C:C,A{i},Extrato!D:D)").number_format = MONEY
    grid(cats, 4, 4 + len(names), 1, 2)
    dv = DataValidation(type="list", formula1=f"=Categorias!$A$5:$A${4 + len(names)}", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"C5:C200")
    widths(ws, {"A": 10, "B": 38, "C": 18, "D": 16})
    widths(cats, {"A": 20, "B": 18})
    wb.save(OUT / "v04-pessoal-x-empresa.xlsx")


# ------------------------------------------------------------------ V08
def v08():
    wb = Workbook()
    ws = wb.active
    ws.title = "Agenda"
    title(ws, "V08 · Agenda financeira com saldo dia a dia")
    label(ws, "A4", "Saldo inicial")
    ws["D4"] = 6000
    ws["D4"].number_format = MONEY
    header(ws, 6, ["Dia", "Descrição", "Tipo", "Valor", "Saldo"])
    rows = [
        (2, "Aluguel", "A pagar", -4500),
        (5, "Fornecedor Alfa", "A pagar", -6200),
        (8, "Vendas no cartão", "A receber", 7400),
        (10, "Fornecedor Beta", "A pagar", -3900),
        (12, "Cliente B", "A receber", 8300),
        (15, "Folha de pagamento", "A pagar", -9800),
        (18, "Vendas", "A receber", 6500),
        (20, "Guia DAS", "A pagar", -2100),
        (26, "Vendas", "A receber", 5900),
    ]
    for i, r in enumerate(rows, 7):
        for j, v in enumerate(r, 1):
            ws.cell(row=i, column=j, value=v)
        ws.cell(row=i, column=4).number_format = MONEY
        prev = "$D$4" if i == 7 else f"E{i - 1}"
        ws.cell(row=i, column=5, value=f"={prev}+D{i}").number_format = MONEY
    last = 6 + len(rows)
    grid(ws, 6, last, 1, 5)
    neg_red(ws, f"E7:E{last}")
    label(ws, f"B{last + 2}", "Saldo final do mês")
    ws[f"E{last + 2}"] = f"=E{last}"
    label(ws, f"B{last + 3}", "Menor saldo do mês")
    ws[f"E{last + 3}"] = f"=MIN(E7:E{last})"
    label(ws, f"B{last + 4}", "Lançamentos com saldo negativo")
    ws[f"E{last + 4}"] = f'=COUNTIF(E7:E{last},"<0")'
    for r in (last + 2, last + 3):
        ws[f"E{r}"].number_format = MONEY
    dv = DataValidation(type="list", formula1='"A pagar,A receber"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("C7:C200")
    ws["A" + str(last + 6)] = "Dica: mantenha as linhas em ordem de data. Para testar uma renegociação, mude o dia e reordene."
    ws["A" + str(last + 6)].font = Font(italic=True, size=9, color="6E7C75")
    widths(ws, {"A": 8, "B": 30, "C": 14, "D": 16, "E": 16})
    wb.save(OUT / "v08-agenda-financeira.xlsx")


# ------------------------------------------------------------------ V12
def v12():
    wb = Workbook()
    ws = wb.active
    ws.title = "Cenários"
    title(ws, "V12 · Projeção de caixa em dois cenários")
    label(ws, "A4", "Saldo inicial")
    ws["C4"] = 20000
    label(ws, "A5", "Compromissos fixos por mês")
    ws["C5"] = 36000
    label(ws, "A6", "Queda nas vendas (cenário fraco)")
    ws["C6"] = 0.25
    ws["C4"].number_format = MONEY
    ws["C5"].number_format = MONEY
    ws["C6"].number_format = "0%"
    for c in ("C4", "C5", "C6"):
        ws[c].fill = PatternFill("solid", fgColor="FBF6E9")
    header(ws, 8, ["Mês", "Receita esperada", "Receita (cenário fraco)", "Compromissos", "Saldo esperado", "Saldo (cenário fraco)"])
    months = [("Jul", 40000), ("Ago", 38000), ("Set", 42000), ("Out", 45000), ("Nov", 48000), ("Dez", 55000)]
    for i, (m, rec) in enumerate(months, 9):
        ws.cell(row=i, column=1, value=m)
        ws.cell(row=i, column=2, value=rec)
        ws.cell(row=i, column=3, value=f"=B{i}*(1-$C$6)")
        ws.cell(row=i, column=4, value="=-$C$5")
        prev_e = "$C$4" if i == 9 else f"E{i - 1}"
        prev_f = "$C$4" if i == 9 else f"F{i - 1}"
        ws.cell(row=i, column=5, value=f"={prev_e}+B{i}+D{i}")
        ws.cell(row=i, column=6, value=f"={prev_f}+C{i}+D{i}")
        for c in range(2, 7):
            ws.cell(row=i, column=c).number_format = MONEY
    last = 8 + len(months)
    grid(ws, 8, last, 1, 6)
    neg_red(ws, f"E9:F{last}")
    label(ws, f"A{last + 2}", "Menor saldo no cenário fraco")
    ws[f"F{last + 2}"] = f"=MIN(F9:F{last})"
    ws[f"F{last + 2}"].number_format = MONEY
    label(ws, f"A{last + 3}", "Reserva para não ficar negativo")
    ws[f"F{last + 3}"] = f"=MAX(0,-F{last + 2})"
    ws[f"F{last + 3}"].number_format = MONEY
    ws[f"A{last + 5}"] = "Altere as células amarelas (premissas) para testar outros cenários."
    ws[f"A{last + 5}"].font = Font(italic=True, size=9, color="6E7C75")
    widths(ws, {"A": 34, "B": 18, "C": 22, "D": 16, "E": 18, "F": 22})
    wb.save(OUT / "v12-cenarios-de-caixa.xlsx")


# ------------------------------------------------------------------ V15
def v15():
    wb = Workbook()
    ws = wb.active
    ws.title = "Conciliação"
    title(ws, "V15 · Conciliação bancária: registros × extrato")
    header(ws, 4, ["Data", "Registro interno", "Valor registrado", "Lançamento no extrato", "Valor no banco", "Diferença", "Status"])
    rows = [
        ("03/06", "Cliente A", 4200, "Pix Cliente A", 4200),
        ("05/06", "Fornecedor Alfa", -3480, "Boleto Alfa", -3480),
        ("08/06", "Aluguel", -4500, "Aluguel", -4500),
        ("12/06", "Cliente B", 2750, "Pix Cliente B", 2600),
        ("15/06", "Energia", -684, "Energia", -684),
    ]
    for i, r in enumerate(rows, 5):
        for j, v in enumerate(r, 1):
            ws.cell(row=i, column=j, value=v)
        ws.cell(row=i, column=6, value=f"=E{i}-C{i}")
        ws.cell(row=i, column=7, value=f'=IF(ABS(F{i})<0.01,"Conferido","Investigar")')
        for c in (3, 5, 6):
            ws.cell(row=i, column=c).number_format = MONEY
    last = 4 + len(rows)
    grid(ws, 4, last, 1, 7)
    ws.conditional_formatting.add(
        f"G5:G{last}", CellIsRule(operator="equal", formula=['"Investigar"'], fill=PatternFill("solid", fgColor=CORAL_BG))
    )
    label(ws, f"D{last + 2}", "Diferença total")
    ws[f"F{last + 2}"] = f"=SUM(F5:F{last})"
    ws[f"F{last + 2}"].number_format = MONEY
    label(ws, f"D{last + 3}", "Itens a investigar")
    ws[f"F{last + 3}"] = f'=COUNTIF(G5:G{last},"Investigar")'
    ws[f"A{last + 5}"] = "Causas comuns de diferença: desconto não registrado, taxa ou tarifa, erro de digitação, pagamento parcial."
    ws[f"A{last + 5}"].font = Font(italic=True, size=9, color="6E7C75")
    widths(ws, {"A": 9, "B": 22, "C": 17, "D": 24, "E": 16, "F": 14, "G": 12})
    wb.save(OUT / "v15-conciliacao-bancaria.xlsx")


# ------------------------------------------------------------------ V23
def v23():
    wb = Workbook()
    ws = wb.active
    ws.title = "Resultado"
    title(ws, "V23 · Fechamento gerencial do mês (maio)")
    header(ws, 4, ["Grupo", "Item", "Valor"])
    rows = [
        ("Receitas", "Serviços prestados", 48000),
        ("Receitas", "Venda de produtos", 12500),
        ("Custos", "Insumos e materiais", -14800),
        ("Pessoal", "Folha e encargos", -21300),
        ("Despesas fixas", "Aluguel e condomínio", -7200),
        ("Operacionais", "Sistemas, energia, internet", -3400),
        ("Tributos", "Impostos sobre vendas", -5445),
    ]
    for i, r in enumerate(rows, 5):
        for j, v in enumerate(r, 1):
            ws.cell(row=i, column=j, value=v)
        ws.cell(row=i, column=3).number_format = MONEY
    last = 4 + len(rows)
    grid(ws, 4, last, 1, 3)
    label(ws, f"B{last + 2}", "Receitas do mês")
    ws[f"C{last + 2}"] = f'=SUMIF(A5:A{last},"Receitas",C5:C{last})'
    label(ws, f"B{last + 3}", "Resultado do mês")
    ws[f"C{last + 3}"] = f"=SUM(C5:C{last})"
    label(ws, f"B{last + 4}", "Margem")
    ws[f"C{last + 4}"] = f"=IF(C{last + 2}=0,0,C{last + 3}/C{last + 2})"
    ws[f"C{last + 2}"].number_format = MONEY
    ws[f"C{last + 3}"].number_format = MONEY
    ws[f"C{last + 4}"].number_format = PCT
    widths(ws, {"A": 18, "B": 32, "C": 16})

    pend = wb.create_sheet("Pendências")
    title(pend, "Contas em aberto")
    header(pend, 4, ["Vencimento", "Descrição", "Valor"])
    for i, r in enumerate([("05/06", "Cliente C", 6800), ("08/06", "Fornecedor Alfa", -4900), ("20/06", "Impostos do mês", -5445)], 5):
        for j, v in enumerate(r, 1):
            pend.cell(row=i, column=j, value=v)
        pend.cell(row=i, column=3).number_format = MONEY
    grid(pend, 4, 7, 1, 3)
    label(pend, "B9", "Saldo das pendências")
    pend["C9"] = "=SUM(C5:C7)"
    pend["C9"].number_format = MONEY
    widths(pend, {"A": 13, "B": 28, "C": 16})

    proj = wb.create_sheet("Projeção")
    title(proj, "Caixa projetado (90 dias)")
    label(proj, "A4", "Saldo atual")
    proj["C4"] = 24600
    proj["C4"].number_format = MONEY
    header(proj, 6, ["Período", "Entradas previstas", "Saídas previstas", "Saldo projetado"])
    for i, (p, e, s) in enumerate([("30 dias", 58000, -56500), ("60 dias", 54000, -57300), ("90 dias", 63000, -56400)], 7):
        proj.cell(row=i, column=1, value=p)
        proj.cell(row=i, column=2, value=e)
        proj.cell(row=i, column=3, value=s)
        prev = "$C$4" if i == 7 else f"D{i - 1}"
        proj.cell(row=i, column=4, value=f"={prev}+B{i}+C{i}")
        for c in (2, 3, 4):
            proj.cell(row=i, column=c).number_format = MONEY
    grid(proj, 6, 9, 1, 4)
    neg_red(proj, "D7:D9")
    widths(proj, {"A": 14, "B": 20, "C": 18, "D": 18})
    wb.save(OUT / "v23-fechamento-do-mes.xlsx")


if __name__ == "__main__":
    for fn in (v04, v08, v12, v15, v23):
        fn()
    print("ok:", ", ".join(sorted(p.name for p in OUT.glob("*.xlsx"))))
